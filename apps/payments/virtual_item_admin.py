"""虚拟道具后台工具：导出微信导入 Excel + 一键同步绑定。

被 Django admin 的 action/button 调用。
"""
import os

from django.conf import settings

from apps.catalog.virtual_item_naming import generate_virtual_item_id

from .models import VirtualProductBinding


def collect_unbound_items(packages=None):
    """收集未绑定微信道具的商品规格。

    返回 [(package, spec, item_id, name, price)]，已绑定的规格跳过。
    """
    from apps.catalog.models import Package, PackageSpec

    bound_spec_ids = set(
        VirtualProductBinding.objects.filter(is_active=True, spec__isnull=False)
        .values_list('spec_id', flat=True)
    )
    bound_pkg_ids = set(
        VirtualProductBinding.objects.filter(is_active=True, package__isnull=False, spec__isnull=True)
        .values_list('package_id', flat=True)
    )

    queryset = Package.objects.filter(is_active=True)
    if packages is not None:
        queryset = queryset.filter(id__in=[p.pk for p in packages])

    results = []
    for pkg in queryset.prefetch_related('specs'):
        if pkg.id in bound_pkg_ids:
            continue
        for spec in pkg.specs.filter(is_active=True).order_by('sort_order', 'id'):
            if spec.id in bound_spec_ids:
                continue
            item_id, name, price = generate_virtual_item_id(pkg, spec)
            if item_id:
                results.append((pkg, spec, item_id, name, price))
    return results


def build_import_excel(packages=None):
    """为未绑定商品规格生成微信批量导入 Excel 文件。

    返回 (文件路径, 条目数)。文件放在 media/wechat_thumb/ 下。
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from PIL import Image

    from apps.catalog.models import Package, PackageSpec

    items = collect_unbound_items(packages)
    if not items:
        return None, 0

    host = 'https://api.huc125.cn'
    thumb_dir = os.path.join(settings.MEDIA_ROOT, 'wechat_thumb')
    os.makedirs(thumb_dir, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'

    hfont = Font(bold=True, size=11)
    hfill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    halign = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    for col, h in enumerate(['道具id', '道具名称', '道具图片', '道具价格', '备注'], 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hfont
        c.fill = hfill
        c.alignment = halign
        c.border = border

    for i, (pkg, spec, item_id, name, price) in enumerate(items, 2):
        # 生成 200x200 缩略图（优先用商品封面，无则生成占位图）
        thumb_url = _ensure_thumbnail(thumb_dir, item_id, pkg, host)
        ws.cell(row=i, column=1, value=item_id).border = border
        ws.cell(row=i, column=2, value=name).border = border
        c3 = ws.cell(row=i, column=3, value=thumb_url)
        c3.border = border
        c3.alignment = Alignment(wrap_text=True)
        ws.cell(row=i, column=4, value=int(price)).border = border
        ws.cell(row=i, column=5, value=f'{pkg.name} {spec.name}').border = border

    out = os.path.join(thumb_dir, 'wechat_import_items.xlsx')
    wb.save(out)
    return out, len(items)


def _ensure_thumbnail(thumb_dir, item_id, pkg, host):
    """生成或复用 200x200 缩略图，返回可访问 URL。"""
    from PIL import Image

    target = os.path.join(thumb_dir, f'{item_id}.jpg')
    if os.path.exists(target):
        return f'{host}/media/wechat_thumb/{item_id}.jpg'

    src = None
    for field in ('cover_url', 'image_url'):
        raw = getattr(pkg, field, '') or ''
        if raw:
            path = raw
            if path.startswith('/media/'):
                path = os.path.join(settings.MEDIA_ROOT, path[len('/media/'):])
            elif path.startswith('http'):
                path = path.split('/media/')[-1]
                path = os.path.join(settings.MEDIA_ROOT, path)
            if os.path.exists(path):
                src = path
                break

    if src:
        img = Image.open(src)
        size = min(img.width, img.height)
        left = (img.width - size) // 2
        top = (img.height - size) // 2
        img = img.crop((left, top, left + size, top + size))
        img = img.resize((200, 200), Image.LANCZOS)
        if img.mode != 'RGB':
            img = img.convert('RGB')
        img.save(target, 'JPEG', quality=75, optimize=True)
    else:
        img = Image.new('RGB', (200, 200), (47, 155, 99))
        img.save(target, 'JPEG', quality=75, optimize=True)
    return f'{host}/media/wechat_thumb/{item_id}.jpg'


def sync_bindings_from_wechat_items(packages=None, dry_run=False):
    """一键同步：为未绑定商品规格创建 VirtualProductBinding。

    规则：按 generate_virtual_item_id 生成道具 ID；已绑定跳过；ID 冲突跳过。
    返回统计 dict。
    """
    from apps.catalog.models import Package, PackageSpec

    items = collect_unbound_items(packages)
    existing_ids = set(
        VirtualProductBinding.objects.filter(is_active=True)
        .values_list('product_id', flat=True)
    )
    created = skipped_duplicate = skipped_invalid = 0
    for pkg, spec, item_id, name, price in items:
        if item_id in existing_ids:
            skipped_duplicate += 1
            continue
        price_fen = int(round(float(price) * 100))
        if price_fen <= 0:
            skipped_invalid += 1
            continue
        if dry_run:
            created += 1
            continue
        VirtualProductBinding.objects.create(
            package=pkg,
            spec=spec,
            product_id=item_id,
            goods_price_fen=price_fen,
            is_active=True,
            remark=f'{pkg.name} {spec.name}',
        )
        existing_ids.add(item_id)
        created += 1
    return {
        'created': created,
        'skipped_duplicate': skipped_duplicate,
        'skipped_invalid': skipped_invalid,
    }
